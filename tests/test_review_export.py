from io import BytesIO
from openpyxl import load_workbook

from screener.review_export import review_excel


def test_export_preserves_snapshot_types_and_all_outcomes():
    run = dict(id='test', target='2026-10-26', status='done', total=3,
               created_at='2026-10-26T16:00:00+08:00', updated_at='2026-10-26T17:00:00+08:00',
               config=dict(source='tencent', boards=['深市主板'], exclude_st=True))
    rows = [dict(code='000001', name='=unsafe', outcome='match', close=12.3,
                 above_days=3, above_days_at_least=True, reason='持续站上'),
            dict(code='000002', name='样本', outcome='no_match', reason='未站上'),
            dict(code='000003', name='缺数据', outcome='unknown', reason='目标日的日K尚未返回')]
    wb = load_workbook(BytesIO(review_excel(run, rows)))
    assert wb.sheetnames == ['复盘概要', '符合个股', '全部明细']
    sheet = wb['全部明细']
    values = list(sheet.values)
    records = [dict(zip(values[0], r)) for r in values[1:]]
    assert len(records) == 3 and wb['符合个股'].max_row == 2
    assert records[0]['代码'] == '000001'
    assert records[0]['日K收盘价'] == 12.3
    assert records[0]['站上天数为下限'] == '是'
    assert records[2]['结果'] == '无法判定'
    assert records[2]['日K收盘价'] is None
    assert not any(c.data_type == 'f' for row in sheet for c in row)
    assert sheet.freeze_panes == 'C2' and sheet.auto_filter.ref


def test_empty_matches_still_has_headers():
    wb = load_workbook(BytesIO(review_excel(dict(config={}, total=0), [])))
    assert wb['符合个股'].max_row == 1


def test_completed_review_shows_download_button(tmp_path, monkeypatch):
    from pathlib import Path
    from streamlit.testing.v1 import AppTest
    from screener.demo import seed
    from screener.store import Store
    from screener.review import ReviewRepository, ReviewRunner
    from screener.provider import Provider
    monkeypatch.setenv('SCREENER_DATA_DIR', str(tmp_path))
    store = Store(tmp_path / 'market.db')
    seed(store)
    repo = ReviewRepository(store)
    runner = ReviewRunner(repo)
    run = runner.create(dict(boards=['沪市主板'], source='tencent', requested_date='2026-10-19',
                            local_only=True, exclude_st=False, exclude_new=False))
    runner.run(run['id'])
    monkeypatch.setattr(Provider, 'call', lambda *a, **k: (_ for _ in ()).throw(AssertionError('network')))
    app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / 'app.py')).run(timeout=30)
    app.radio(key='workspace_page').set_value('每日复盘').run(timeout=30)
    assert not app.exception
    assert any(b.label == '导出个股明细 Excel' for b in app.get('download_button'))

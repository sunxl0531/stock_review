"""Export saved review details without requesting market data."""
from collections import Counter
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

OUTCOMES = {'match': '符合', 'no_match': '未符合', 'excluded': '已排除',
            'unknown': '无法判定', 'error': '请求失败'}
FIELDS = {'code': '代码', 'name': '名称', 'board': '市场', 'industry': '行业板块',
          'outcome': '结果', 'reason': '判断原因', 'indicator_date': '指标日期',
          'close': '日K收盘价', 'snakea': 'SNAKEA', 'snakea_previous': '前一日SNAKEA',
          'snakea_direction': 'SNAKEA方向', 'above_days': '连续站上天数',
          'above_days_at_least': '站上天数为下限', 'break_date': '突破日期',
          'distance_pct': '距SNAKEA %', 'snakea_change_pct': 'SNAKEA日变化 %',
          'stop_lower': '停损下限', 'stop_color': '停损显示', 'verification': '数据依据',
          'cache_mode': '历史读取方式', 'data_fetched_at': '行情获取时间', 'fetched_at': '计算时间'}


def _sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'  # Provider strings must not become Excel formulas.
            elif isinstance(cell.value, float):
                cell.number_format = '0.00'
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='176A80')
    ws.freeze_panes = 'C2' if len(headers) > 2 else 'A2'
    ws.auto_filter.ref = ws.dimensions
    for index, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(index)].width = (
            42 if header in ('判断原因', '数据依据', '内容') else 28 if '时间' in header else 20)


def review_excel(run, results):
    wb = Workbook()
    wb.remove(wb.active)
    cfg = run.get('config', {})
    counts = Counter(row.get('outcome') for row in results)
    summary = [('复盘交易日', run.get('target')), ('记录ID', run.get('id')),
               ('状态', run.get('status')), ('开始时间', run.get('created_at')),
               ('最后更新时间', run.get('updated_at')), ('目标股票数', run.get('total', 0)),
               ('已处理股票数', len(results)),
               *[(label, counts[key]) for key, label in OUTCOMES.items()],
               ('市场', '、'.join(cfg.get('boards', []))),
               ('历史来源', cfg.get('source')), ('排除ST', '是' if cfg.get('exclude_st') else '否'),
               ('排除上市不足60个交易日', '是' if cfg.get('exclude_new') else '否'),
               ('说明', '本次复盘保存的结果；缺失指标留空。百分比列以百分数值表示。')]
    _sheet(wb, '复盘概要', ['项目', '内容'], summary)
    ordered = sorted(results, key=lambda r: (r.get('board') or '', r.get('code') or ''))
    def values(row):
        data = row | {'outcome': OUTCOMES.get(row.get('outcome'), row.get('outcome')),
                      'above_days_at_least': ('是' if row.get('above_days_at_least') else '否')
                      if row.get('above_days') is not None else None}
        return [data.get(key) for key in FIELDS]
    _sheet(wb, '符合个股', list(FIELDS.values()), [values(r) for r in ordered if r.get('outcome') == 'match'])
    _sheet(wb, '全部明细', list(FIELDS.values()), [values(r) for r in ordered])
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
